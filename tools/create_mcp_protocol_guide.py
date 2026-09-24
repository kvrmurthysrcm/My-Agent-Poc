from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("artifacts/MCP_Protocol_and_Project_Usage_Guide.docx")


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(color)
    run_pr.append(underline)
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break()
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 32, 32)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)

    for style_name, size in (("Heading 1", 17), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)


def add_feature_table(doc):
    rows = [
        ("Tools", "Execute typed operations", "Searches, API calls, database operations, and controlled mutations", "Used"),
        ("Resources", "Read URI-addressable context", "Documents, schemas, records, and other reusable context", "Not used"),
        ("Prompts", "Publish reusable interaction templates", "Shared workflows, few-shot examples, and guided tasks", "Not used"),
        ("Elicitation", "Request user input or confirmation", "Missing parameters, ambiguity, approvals, and secure redirects", "Not used"),
        ("Notifications", "Report changes without polling", "Dynamic tool catalogs and changing resources", "Not used"),
        ("Progress and cancellation", "Manage an active long operation", "Indexing, exports, migrations, and batch processing", "Not used"),
        ("Tasks", "Represent durable asynchronous work", "Jobs that continue after the initial request", "Not used"),
        ("Completion", "Suggest argument values", "Prompt parameters, resource paths, names, and IDs", "Not used"),
        ("Authentication", "Protect remote MCP access", "External or multi-user MCP deployments", "Gateway only"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.25, 1.55, 2.75, 0.85]
    headers = ["Capability", "Purpose", "When it is useful", "Project"]
    for idx, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[idx])
        set_cell_fill(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[col_index])
            set_cell_fill(cell, "F5F8FB" if row_index % 2 == 0 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index == 3 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            run.font.size = Pt(8.5)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    title = doc.add_paragraph(style="Title")
    title.add_run("MCP Protocol and Project Usage Guide")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Online Library MCP in My Agent POC")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(70, 70, 70)

    p = doc.add_paragraph()
    p.add_run("Main conclusion  ").bold = True
    p.add_run(
        "The Online Library MCP service uses a focused part of the Model Context Protocol. "
        "It exposes 24 typed tools over stateless Streamable HTTP, while the Online Library Agent "
        "handles natural-language interpretation and tool selection. Resources, prompts, elicitation, "
        "notifications, progress reporting, and asynchronous tasks are not currently implemented."
    )

    doc.add_heading("What MCP Provides", level=1)
    doc.add_paragraph(
        "The Model Context Protocol is a standard contract between an AI host and an external capability server. "
        "It defines how a client discovers capabilities, sends structured requests, and receives results. MCP does "
        "not decide which tool should answer a question, implement business rules, or query PostgreSQL by itself."
    )
    add_code_block(
        doc,
        [
            "User",
            "  -> AI host and agent",
            "      -> MCP client",
            "          -> MCP server",
            "              -> Online Library API",
            "                  -> PostgreSQL",
        ],
    )

    doc.add_heading("Protocol Layers", level=2)
    add_bullet(doc, "The data layer uses JSON-RPC messages for discovery, capability definitions, requests, responses, errors, and notifications.")
    add_bullet(doc, "The transport layer carries those messages through local standard input and output or remote Streamable HTTP.")
    add_bullet(doc, "The SDK handles message framing, protocol metadata, capability negotiation, and validation so application code can focus on business functions.")

    doc.add_heading("MCP Capability Families", level=1)
    doc.add_paragraph(
        "Tools are only one MCP primitive. The protocol also provides ways to publish context, distribute reusable prompts, "
        "collect user input, announce changes, and coordinate long-running work."
    )
    add_feature_table(doc)

    doc.add_heading("How Tools Work", level=1)
    doc.add_paragraph(
        "A tool is an executable function described by a stable name, a human-readable description, and a JSON Schema for its arguments. "
        "The client discovers tools through tools/list and invokes one through tools/call. Results can contain text, structured data, "
        "images, audio, resource links, or embedded resources."
    )
    doc.add_heading("Tool flow in this project", level=2)
    steps = [
        "The Online Library Agent opens a Streamable HTTP connection to the MCP endpoint at /mcp.",
        "The Python MCP client initializes the connection and requests tools/list.",
        "FastMCP returns the registered tool names, descriptions, defaults, and input schemas.",
        "The OnlineLibraryPlanner applies domain guardrails and business overrides, then asks Ollama to select a tool when needed.",
        "The client sends tools/call with the selected name and validated arguments.",
        "The MCP tool calls a fixed Online Library REST endpoint with httpx.",
        "The REST API executes parameterized SQL against PostgreSQL and returns JSON.",
        "FastMCP packages the result as an MCP tool response, and the Agent summarizes it for the user.",
    ]
    for step in steps:
        add_number(doc, step)

    doc.add_heading("What the Project Exposes", level=1)
    doc.add_paragraph(
        "The server registers 24 read-only tools. They cover catalog search and resource details; books by author, genre, or tag; "
        "authors and facets; users, subscriptions, and approvals; reading progress and bookshelves; subscription tiers and rules; "
        "and allowlisted table retrieval for diagnostics."
    )
    doc.add_paragraph(
        "FastMCP derives each tool definition from the decorated Python function, its type annotations, its defaults, and its docstring. "
        "The MCP module owns this protocol-facing contract and the mapping to REST endpoints. The Online Library API remains responsible "
        "for SQL and business data access."
    )

    doc.add_heading("What MCP Does Not Do in This Project", level=2)
    add_bullet(doc, "MCP does not convert the user's NLQ into SQL.")
    add_bullet(doc, "MCP does not choose the tool. The OnlineLibraryPlanner performs that work.")
    add_bullet(doc, "MCP does not query PostgreSQL directly. The Online Library API owns the repository code.")
    add_bullet(doc, "MCP does not generate the final natural-language answer. The Agent summarizes the tool result.")
    add_bullet(doc, "MCP does not cache business data. The API and PostgreSQL remain the sources of truth.")

    doc.add_heading("When the Other Features Become Useful", level=1)
    doc.add_heading("Resources", level=2)
    doc.add_paragraph(
        "Resources expose read-only context through stable URIs. They are preferable to tools when clients should browse or attach data "
        "without framing access as an action. Useful project examples include library://schema/catalog, library://resources/{id}, and "
        "library://subscription-rules. Resource templates can describe parameterized URI patterns, and subscriptions can report changes."
    )

    doc.add_heading("Prompts and completion", level=2)
    doc.add_paragraph(
        "Prompts publish reusable interaction templates that clients can discover and present to users. They would be useful if several "
        "clients needed the same workflows for finding books, reviewing pending users, or explaining subscription rules. Completion can "
        "suggest valid prompt arguments or resource identifiers while a user types. The current Agent owns its guided questions and prompts, "
        "so adding MCP prompts now would duplicate that logic."
    )

    doc.add_heading("Elicitation", level=2)
    doc.add_paragraph(
        "Elicitation lets the server request missing information or confirmation from the user. It becomes useful when a book title is ambiguous, "
        "a required resource ID is missing, or a future write operation needs explicit approval. Current tools return validation errors instead, "
        "which is adequate for short, read-only calls."
    )

    doc.add_heading("Notifications", level=2)
    doc.add_paragraph(
        "Notifications keep clients synchronized when a tool catalog or resource changes. They matter when tools are enabled dynamically by role, "
        "when data sources appear or disappear, or when clients subscribe to changing resources. The current tools are registered statically at startup, "
        "and the server is configured for stateless JSON responses, so notification streams provide little value today."
    )

    doc.add_heading("Progress cancellation and tasks", level=2)
    doc.add_paragraph(
        "Progress and cancellation apply to active long-running requests. Tasks represent durable asynchronous work that a client can check later. "
        "These capabilities would fit document ingestion, vector reindexing, bulk imports, graph construction, large exports, and administrative jobs. "
        "They are unnecessary for the current tools because each tool performs a short, read-only HTTP GET."
    )

    doc.add_heading("Sampling roots and logging", level=2)
    doc.add_paragraph(
        "Older MCP revisions included sampling so a server could request an LLM completion from its host. The current protocol deprecates sampling, "
        "and this project already calls Ollama from the Agent. Roots are relevant to local filesystem servers that need an allowed directory scope; "
        "they do not apply to this HTTP API wrapper. Protocol-level logging is also unnecessary here because the services use normal application logs "
        "and forwarded trace identifiers."
    )

    doc.add_heading("Transport Security and Health", level=1)
    doc.add_paragraph(
        "The MCP server runs on port 8004 and exposes stateless Streamable HTTP at /mcp. Kubernetes currently uses TCP socket probes for liveness and "
        "readiness. The Secure API performs a deeper functional check by sending tools/list. The direct MCP server does not currently configure MCP-native "
        "OAuth authorization, so network controls and the authenticated gateway form the security boundary. A production remote MCP service should add "
        "endpoint authentication and a readiness check that verifies the downstream Online Library API."
    )

    doc.add_heading("Project Usage Summary", level=1)
    summary = [
        ("Used", "Streamable HTTP, JSON-RPC request and response handling, SDK lifecycle initialization, tools/list, tools/call, generated input schemas, and JSON-shaped tool results."),
        ("SDK managed", "Protocol framing, capability metadata, validation, and standard MCP response packaging."),
        ("Not used", "Resources, resource templates, prompts, completion, elicitation, subscriptions, progress, durable tasks, sampling, roots, and rich media results."),
        ("Application owned", "Domain guardrails, NLQ interpretation, tool selection, REST calls, SQL execution, authorization at the gateway, and final answer generation."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col_index, (cell, value) in enumerate(zip(table.rows[0].cells, ("Ownership", "Current implementation"))):
        cell.width = Inches(1.35 if col_index == 0 else 5.05)
        set_cell_fill(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell, top=130, bottom=130)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(summary, start=1):
        cells = table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(1.35 if col_index == 0 else 5.05)
            set_cell_fill(cell, "EAF1F7" if col_index == 0 else ("F8FAFC" if row_index % 2 else "FFFFFF"))
            set_cell_border(cell)
            set_cell_margins(cell, top=130, bottom=130)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            run = cell.paragraphs[0].add_run(value)
            run.bold = col_index == 0
            run.font.size = Pt(9)

    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(
        "The current scope is appropriate for this POC. Keep tools as the primary interface while operations remain short and read-only. Add resources "
        "when clients need addressable reusable context, elicitation when tools need clarification or confirmation, and tasks with progress when the MCP "
        "surface begins to expose ingestion or administrative jobs. Before external deployment, add MCP endpoint authentication and dependency-aware readiness."
    )

    doc.add_heading("Sources", level=1)
    sources = [
        ("MCP architecture overview", "https://modelcontextprotocol.io/docs/2026-07-28/learn/architecture"),
        ("MCP server concepts", "https://modelcontextprotocol.io/docs/2026-07-28/learn/server-concepts"),
        ("MCP Tasks extension", "https://tasks.extensions.modelcontextprotocol.io/specification/draft/tasks"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, label, url)
    p = doc.add_paragraph(style="List Bullet")
    p.add_run("Project implementation: ").bold = True
    p.add_run("modules/online_library_mcp, modules/online_library_agent, modules/online_library, and modules/secure_api")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("My Agent POC  |  MCP Protocol Guide")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(95, 95, 95)

    doc.core_properties.title = "MCP Protocol and Project Usage Guide"
    doc.core_properties.subject = "Model Context Protocol capabilities and their use in My Agent POC"
    doc.core_properties.author = "My Agent POC"
    doc.core_properties.keywords = "MCP, Model Context Protocol, FastMCP, tools, resources, prompts"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
