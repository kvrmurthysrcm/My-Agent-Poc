from html.parser import HTMLParser
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from app.core.config import Settings, get_settings
from app.services.pdf_parsers.factory import PdfParserFactory
from app.services.pdf_text_cleanup import PdfTextCleanupService
from app.services.text_extraction_types import ExtractionResult


class TextExtractionService:
    def __init__(self, settings: Settings | None = None):
        self.settings = settings or get_settings()

    def extract(self, path: Path, extension: str) -> ExtractionResult:
        if extension == ".txt":
            return self._extract_txt(path)
        if extension == ".pdf":
            return self._extract_pdf(path)
        if extension == ".docx":
            return self._extract_docx(path)
        if extension == ".epub":
            return self._extract_epub(path)
        raise ValueError(f"Unsupported extraction extension: {extension}")

    def _extract_txt(self, path: Path) -> ExtractionResult:
        text = path.read_text(encoding="utf-8", errors="replace")
        return ExtractionResult(text=text, parser_name="plain_text", page_count=None, metadata={})

    def _extract_pdf(self, path: Path) -> ExtractionResult:
        parser = PdfParserFactory.build(self.settings)
        try:
            return self._clean_pdf_result(parser.extract(path))
        except RuntimeError:
            if self.settings.pdf_parser == "pymupdf":
                result = PdfParserFactory.build(self.settings.model_copy(update={"pdf_parser": "pypdf"})).extract(path)
                return self._clean_pdf_result(result)
            raise

    def _clean_pdf_result(self, result: ExtractionResult) -> ExtractionResult:
        cleaned_text, stats = PdfTextCleanupService(self.settings).clean(result.text)
        return ExtractionResult(
            text=cleaned_text,
            parser_name=result.parser_name,
            page_count=result.page_count,
            metadata={
                **result.metadata,
                "pdf_text_cleanup": {
                    "drop_caps_repaired": self.settings.pdf_repair_drop_caps,
                    "repeated_headers_footers_removed": self.settings.pdf_remove_repeated_headers_footers,
                    "dehyphenated_line_breaks": self.settings.pdf_dehyphenate_line_breaks,
                    "private_use_glyphs_removed": self.settings.pdf_remove_private_use_glyphs,
                    "unicode_normalized": self.settings.pdf_normalize_unicode,
                    "boilerplate_lines_removed": self.settings.pdf_remove_boilerplate_lines,
                    "page_number_lines_removed": self.settings.pdf_remove_page_number_lines,
                    "joined_words_repaired": self.settings.pdf_repair_joined_words,
                    "repeated_line_count": stats.repeated_lines_removed,
                    "boilerplate_line_count": stats.boilerplate_lines_removed,
                    "page_number_line_count": stats.page_number_lines_removed,
                    "joined_word_repair_count": stats.joined_word_repairs,
                },
            },
        )

    def _extract_docx(self, path: Path) -> ExtractionResult:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required for DOCX extraction") from exc

        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return ExtractionResult(text="\n\n".join(parts), parser_name="python-docx", page_count=None, metadata={})

    def _extract_epub(self, path: Path) -> ExtractionResult:
        html_files: list[str] = []
        metadata: dict = {}
        with ZipFile(path) as archive:
            rootfile = self._find_epub_rootfile(archive)
            if rootfile:
                html_files = self._find_epub_spine_files(archive, rootfile)
                metadata["opf_path"] = rootfile
            if not html_files:
                html_files = [
                    name
                    for name in archive.namelist()
                    if name.lower().endswith((".xhtml", ".html", ".htm"))
                    and not name.lower().startswith("meta-inf/")
                ]

            parts = []
            for name in html_files:
                try:
                    raw_html = archive.read(name).decode("utf-8", errors="replace")
                except KeyError:
                    continue
                text = _HtmlTextExtractor.extract(raw_html)
                if text:
                    parts.append(text)

        return ExtractionResult(
            text="\n\n".join(parts),
            parser_name="epub-zip-html",
            page_count=None,
            metadata={**metadata, "content_files": html_files},
        )

    def _find_epub_rootfile(self, archive: ZipFile) -> str | None:
        try:
            container_xml = archive.read("META-INF/container.xml")
        except KeyError:
            return None
        root = ET.fromstring(container_xml)
        for element in root.iter():
            if element.tag.endswith("rootfile"):
                return element.attrib.get("full-path")
        return None

    def _find_epub_spine_files(self, archive: ZipFile, rootfile: str) -> list[str]:
        try:
            root = ET.fromstring(archive.read(rootfile))
        except (KeyError, ET.ParseError):
            return []

        manifest: dict[str, str] = {}
        spine_ids: list[str] = []
        for element in root.iter():
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "item" and element.attrib.get("id") and element.attrib.get("href"):
                manifest[element.attrib["id"]] = element.attrib["href"]
            elif tag == "itemref" and element.attrib.get("idref"):
                spine_ids.append(element.attrib["idref"])

        base_path = str(Path(rootfile).parent).replace("\\", "/")
        if base_path == ".":
            base_path = ""

        files: list[str] = []
        for item_id in spine_ids:
            href = manifest.get(item_id)
            if not href:
                continue
            candidate = f"{base_path}/{href}" if base_path else href
            files.append(candidate.replace("\\", "/"))
        return files


class _HtmlTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    @classmethod
    def extract(cls, html: str) -> str:
        parser = cls()
        parser.feed(html)
        return " ".join(" ".join(parser.parts).split())

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"script", "style", "nav"}:
            self._skip_depth += 1
        if tag.lower() in {"p", "div", "section", "article", "h1", "h2", "h3", "h4", "li", "br"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "nav"} and self._skip_depth:
            self._skip_depth -= 1
        if tag.lower() in {"p", "div", "section", "article", "h1", "h2", "h3", "h4", "li"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0 and data.strip():
            self.parts.append(data.strip())
